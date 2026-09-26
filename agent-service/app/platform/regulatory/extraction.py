"""Text extraction for Regulatory Knowledge Base source files —
specs/platform/10-regulatory-knowledge-base-spec.md, "Phase 2
addendum". Pure: bytes in, normalized text out. The bytes themselves
live in regulatory_source_files (app-api-owned — uploads are too large
for a Temporal payload); the caller reads them from Postgres.

PDF extraction reads the text layer only. There is deliberately no OCR:
a scanned PDF fails with a clear reason rather than producing garbage
that would later be cited as regulatory text."""

from __future__ import annotations

import io
from dataclasses import dataclass
from html.parser import HTMLParser
from typing import ClassVar

from app.platform.regulatory.chunking import normalize_text

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TXT = "text/plain"
HTML = "text/html"
SUPPORTED_CONTENT_TYPES = (PDF, DOCX, TXT, HTML)

_MIN_EXTRACTED_CHARS = 20


class ExtractionError(ValueError):
    """The source can't be turned into usable text — shown to the
    officer as-is (400), never retried."""


@dataclass
class ExtractionResult:
    text: str
    page_count: int | None = None


def extract_text(content: bytes, content_type: str) -> ExtractionResult:
    base_type = content_type.split(";", 1)[0].strip().lower()
    if base_type == PDF:
        result = _extract_pdf(content)
    elif base_type == DOCX:
        result = ExtractionResult(text=_extract_docx(content))
    elif base_type == TXT:
        result = ExtractionResult(text=_decode_text(content))
    elif base_type == HTML:
        result = ExtractionResult(text=_extract_html(_decode_text(content)))
    else:
        raise ExtractionError(f"unsupported content type {content_type!r} — use PDF, DOCX, TXT or HTML")

    result.text = normalize_text(result.text)
    if len(result.text) < _MIN_EXTRACTED_CHARS:
        if base_type == PDF:
            raise ExtractionError(
                "no extractable text in this PDF — it may be a scanned image. OCR isn't supported; "
                "use a text-based PDF or paste the text instead"
            )
        raise ExtractionError("the source contains no usable text")
    return result


def _extract_pdf(content: bytes) -> ExtractionResult:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ExtractionError("this PDF is password-protected — upload an unprotected copy")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise ExtractionError(f"could not read this PDF: {exc}") from exc
    return ExtractionResult(text="\n\n".join(pages), page_count=len(pages))


def _extract_docx(content: bytes) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(content))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise ExtractionError(f"could not read this DOCX: {exc}") from exc

    blocks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(b for b in blocks if b.strip())


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("text is not valid UTF-8 — re-save it as UTF-8 and upload again") from exc


class _HtmlToText(HTMLParser):
    _SKIP: ClassVar[frozenset[str]] = frozenset({"script", "style", "noscript", "head", "nav", "footer"})
    _BLOCK: ClassVar[frozenset[str]] = frozenset(
        {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}
    )

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self.parts.append(data)


def _extract_html(markup: str) -> str:
    parser = _HtmlToText()
    parser.feed(markup)
    parser.close()
    # Collapse intra-line whitespace runs HTML source formatting leaves behind.
    lines = [" ".join(line.split()) for line in "".join(parser.parts).split("\n")]
    return "\n".join(lines)
