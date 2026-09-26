"""specs/platform/10-regulatory-knowledge-base-spec.md, "Phase 2
addendum" — extraction.py is pure (bytes in, text out). Fixtures are
generated in-test so no binary files live in the repo."""

from __future__ import annotations

import io

import pytest

from app.platform.regulatory.extraction import DOCX, HTML, PDF, TXT, ExtractionError, extract_text


def _minimal_pdf(lines: list[str]) -> bytes:
    """A valid single-page PDF with a real text layer (Helvetica)."""
    text_ops = " ".join(f"({line}) Tj 0 -16 Td" for line in lines)
    stream = f"BT /F1 12 Tf 72 720 Td {text_ops} ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def _docx(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Indicator"
    table.rows[0].cells[1].text = "Structuring"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_pdf_text_layer_is_extracted_with_page_count():
    result = extract_text(_minimal_pdf(["7. Reporting of suspicious transactions", "Every reporting entity shall file"]), PDF)
    assert "Reporting of suspicious transactions" in result.text
    assert result.page_count == 1


def test_pdf_without_text_layer_fails_clearly():
    with pytest.raises(ExtractionError, match="OCR"):
        extract_text(_minimal_pdf([]), PDF)


def test_corrupt_pdf_fails_clearly():
    with pytest.raises(ExtractionError):
        extract_text(b"%PDF-1.4 this is not really a pdf", PDF)


def test_docx_paragraphs_and_tables_are_extracted():
    result = extract_text(_docx(["Section 7A Customer due diligence", "Verify beneficial ownership."]), DOCX)
    assert "Section 7A Customer due diligence" in result.text
    assert "Indicator | Structuring" in result.text


def test_txt_utf8_with_bom():
    result = extract_text("﻿Section 34 — tipping-off is prohibited.".encode(), "text/plain; charset=utf-8")
    assert result.text == "Section 34 — tipping-off is prohibited."


def test_txt_invalid_utf8_fails():
    with pytest.raises(ExtractionError, match="UTF-8"):
        extract_text(b"\xff\xfe\xfa not utf8 at all here", TXT)


def test_html_drops_scripts_navigation_and_keeps_blocks():
    markup = """<html><head><title>x</title><script>alert(1)</script></head><body>
    <nav>Home | About</nav><h1>AML/CFT Regulations</h1><p>Regulation 4.   Customer   due diligence.</p>
    <footer>Copyright</footer></body></html>"""
    result = extract_text(markup.encode(), HTML)
    assert "alert" not in result.text
    assert "Home" not in result.text and "Copyright" not in result.text
    assert "AML/CFT Regulations\n\nRegulation 4. Customer due diligence." in result.text


def test_unsupported_type_and_empty_text():
    with pytest.raises(ExtractionError, match="unsupported"):
        extract_text(b"GIF89a", "image/gif")
    with pytest.raises(ExtractionError, match="no usable text"):
        extract_text(b"   \n  ", TXT)
